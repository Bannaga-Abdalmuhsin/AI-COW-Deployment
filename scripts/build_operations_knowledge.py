#!/usr/bin/env python3
"""Build chunked public operations knowledge files from the source workbook."""

from __future__ import annotations

import json
import sys
from datetime import date, datetime, time, timedelta
from pathlib import Path

from openpyxl import load_workbook


FIELDS = [
    "recordNo", "cowId", "bpmRequestId", "planStatus", "movedAt", "reachedAt",
    "fromLocation", "fromLatitude", "fromLongitude", "category", "toLocation",
    "eventName", "subLocation", "toLatitude", "toLongitude", "administrativeRegion",
    "distanceKm", "radiusUnderContractKm", "movementType", "regionFrom", "regionTo",
    "cityDistrict", "vendor", "kpiScope", "slaSeconds", "kpiSeconds", "exceededSeconds",
    "priority", "cancelled", "statusRemarks", "managementComments", "month",
]


def clean(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.isoformat()
    if isinstance(value, timedelta):
        return int(value.total_seconds())
    if isinstance(value, str):
        value = " ".join(value.split())
        return None if not value or value.upper() == "NA" else value
    return value


def main() -> None:
    if len(sys.argv) != 3:
        raise SystemExit("Usage: build_operations_knowledge.py INPUT.xlsx OUTPUT_DIR")
    source = Path(sys.argv[1])
    output_dir = Path(sys.argv[2])
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    records = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        record = {key: clean(value) for key, value in zip(FIELDS, row)}
        if any(value is not None for value in record.values()):
            records.append(record)
    output_dir.mkdir(parents=True, exist_ok=True)
    chunk_size = 400
    chunks = []
    for index, start in enumerate(range(0, len(records), chunk_size), start=1):
        filename = f"records-{index:02d}.json"
        chunk = records[start:start + chunk_size]
        (output_dir / filename).write_text(
            json.dumps({"records": chunk}, ensure_ascii=False, separators=(",", ":")),
            encoding="utf-8",
        )
        chunks.append({"file": filename, "records": len(chunk)})
    manifest = {
        "source": source.name,
        "sheet": sheet.title,
        "recordCount": len(records),
        "fieldCount": len(FIELDS),
        "fields": FIELDS,
        "chunks": chunks,
    }
    (output_dir / "index.json").write_text(
        json.dumps(manifest, ensure_ascii=False, separators=(",", ":")),
        encoding="utf-8",
    )
    print(f"Wrote {len(records):,} records across {len(chunks)} chunks to {output_dir}")


if __name__ == "__main__":
    main()
